"""Mixed-media intake with stable provenance. No OCR or remote transcription.
python media_ingest.py approved.txt transcript.vtt document.pdf handout.docx
Install pypdf/python-docx/httpx from live_requirements.txt in an isolated env.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re
from urllib.parse import urlsplit
import zipfile

def parse_document(path):
    path=Path(path)
    if path.stat().st_size>10_000_000:raise ValueError('upload exceeds 10 MB teaching limit')
    raw=path.read_bytes()
    if len(raw)>10_000_000:raise ValueError('upload exceeds 10 MB teaching limit')
    digest=hashlib.sha256(raw).hexdigest();suffix=path.suffix.lower();sections=[]
    if suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        if len(reader.pages)>100:raise ValueError('page budget exceeded')
        for number,page in enumerate(reader.pages,1):
            text=(page.extract_text() or '').strip()
            if not text:raise ValueError('empty/scanned page requires a separate OCR review path')
            sections.append({'locator':f'page:{number}','text':text})
    elif suffix=='.docx':
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist())>30_000_000:raise ValueError('expanded archive budget exceeded')
        from docx import Document
        document=Document(path)
        for number,p in enumerate(document.paragraphs,1):
            if p.text.strip():sections.append({'locator':f'paragraph:{number}','text':p.text.strip()})
        for number,table in enumerate(document.tables,1):
            sections.append({'locator':f'table:{number}','text':'\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)})
    elif suffix=='.vtt':
        text=raw.decode('utf-8-sig')
        for block in re.split(r'\n\s*\n',text):
            lines=block.splitlines();timing=next((i for i,l in enumerate(lines) if '-->' in l),None)
            if timing is not None:sections.append({'locator':'time:'+lines[timing].strip(),'text':' '.join(lines[timing+1:]).strip()})
    elif suffix in {'.txt','.md'}:
        sections=[{'locator':'text:1','text':raw.decode('utf-8').strip()}]
    else:raise ValueError('unsupported input format')
    if not sections or any(not s['text'] for s in sections):raise ValueError('no usable text; quarantine instead of activate')
    return [{'document':path.name,'digest':digest,**section} for section in sections]

async def fetch_approved(url,allowed_origins):
    """allowed_origins is server/operator configuration, never a model field.
    Production still needs DNS/IP egress enforcement and parser worker limits.
    """
    import httpx
    parsed=urlsplit(url);origin=f'{parsed.scheme}://{parsed.netloc}'
    if parsed.username or parsed.password or parsed.scheme not in {'http','https'} or origin not in allowed_origins:
        raise ValueError('URL origin not approved')
    chunks=[];size=0
    async with httpx.AsyncClient(follow_redirects=False,timeout=10) as client:
        async with client.stream('GET',url) as response:
            response.raise_for_status()
            if response.is_redirect:raise ValueError('redirects require a new origin decision')
            async for data in response.aiter_bytes():
                size+=len(data)
                if size>1_000_000:raise ValueError('website byte budget exceeded')
                chunks.append(data)
    from html.parser import HTMLParser
    class Text(HTMLParser):
        def __init__(self):super().__init__();self.parts=[];self.skip=0
        def handle_starttag(self,tag,attrs):
            if tag in {'script','style'}:self.skip+=1
        def handle_endtag(self,tag):
            if tag in {'script','style'}:self.skip=max(0,self.skip-1)
        def handle_data(self,data):
            if not self.skip and data.strip():self.parts.append(data.strip())
    raw=b''.join(chunks);parser=Text();parser.feed(raw.decode('utf-8',errors='replace'))
    text=' '.join(parser.parts)
    if not text:raise ValueError('no readable website text')
    return {'document':url,'digest':hashlib.sha256(raw).hexdigest(),'locator':'html:text','text':text}

if __name__=='__main__':
    args=argparse.ArgumentParser();args.add_argument('paths',nargs='+');parsed=args.parse_args()
    for path in parsed.paths:
        for row in parse_document(path):print(json.dumps(row,ensure_ascii=False))
