"""Exercise real local transports/parsers/SDKs; no provider or cloud calls."""
import asyncio,hashlib,hmac,importlib,json,os,subprocess,sys,tempfile
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];EXAMPLES=ROOT/'backend/academy/content/engineering_examples'
sys.path.insert(0,str(EXAMPLES))
from fastapi.testclient import TestClient
from voice_gateway import app
os.environ['VOICE_LAB_TOKEN']='local-synthetic-test-token'
with TestClient(app) as client:
 assert client.get('/health').json()['provider_tested'] is False
 try:
  with client.websocket_connect('/voice',headers={'origin':'https://unapproved.example'}):pass
 except Exception:pass
 else:raise AssertionError('wrong origin accepted')
 with client.websocket_connect('/voice',headers={'origin':'http://127.0.0.1:8092'}) as ws:
  ws.send_json({'token':'wrong'})
  assert ws.receive()['type']=='websocket.close'
 with client.websocket_connect('/voice',headers={'origin':'http://127.0.0.1:8092'}) as ws:
  ws.send_json({'token':os.environ['VOICE_LAB_TOKEN']});assert ws.receive_json()['mode']=='echo'
  pcm=b'\x01\x00'*160;ws.send_bytes(pcm)
  assert ws.receive_json()=={'type':'audio_format','rate':16000,'fixture':True}
  assert ws.receive_bytes()==pcm
  ws.send_json({'type':'audio_end'});assert ws.receive_json()['type']=='audio_stopped'
  ws.send_bytes(b'x');assert ws.receive_json()['type']=='error'
from media_ingest import parse_document,fetch_approved
from reportlab.pdfgen.canvas import Canvas
from docx import Document
with tempfile.TemporaryDirectory() as folder:
 folder=Path(folder)
 pdf=folder/'notes.pdf';canvas=Canvas(str(pdf));canvas.drawString(30,750,'Synthetic returns: 30 days.');canvas.save()
 assert parse_document(pdf)[0]['locator']=='page:1'
 assert '30 days' in parse_document(pdf)[0]['text']
 docx=folder/'notes.docx';doc=Document();doc.add_paragraph('Unopened items only.');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Exception';table.cell(0,1).text='Personalized';doc.save(docx)
 parsed=parse_document(docx);assert len(parsed)==2 and parsed[1]['locator']=='table:1'
 vtt=folder/'clip.vtt';vtt.write_text('WEBVTT\n\n00:00:01.000 --> 00:00:04.000\nPersonalized items are excluded.\n')
 assert parse_document(vtt)[0]['locator'].startswith('time:00:00:01')
 blank=folder/'blank.pdf';canvas=Canvas(str(blank));canvas.showPage();canvas.save()
 try:parse_document(blank)
 except ValueError:pass
 else:raise AssertionError('blank scan accepted')
 try:asyncio.run(fetch_approved('http://169.254.169.254/',{'https://approved.example'}))
 except ValueError:pass
 else:raise AssertionError('unapproved origin accepted')
from workspace_connectors import verify_slack,require_review,payload_hash
raw=b'{"event_id":"e1"}';timestamp=1000
signature='v0='+hmac.new(b'fixture-secret',b'v0:1000:'+raw,hashlib.sha256).hexdigest()
assert verify_slack(raw,timestamp,signature,'fixture-secret',now=1001)['event_id']=='e1'
for body,now in [(b'{}',1001),(raw,2000)]:
 try:verify_slack(body,timestamp,signature,'fixture-secret',now=now)
 except ValueError:pass
 else:raise AssertionError('bad or stale event accepted')
payload={'calendar':'synthetic'};review={'owner':'acme','hash':payload_hash(payload),'expires':1100}
require_review(payload,review,'acme',now=1001)
try:require_review({'calendar':'changed'},review,'acme',now=1001)
except PermissionError:pass
else:raise AssertionError('changed effect accepted')
# Actual installed SDK contracts, without any provider API request.
import inspect
from google import genai
from google.genai import types
c=genai.Client(api_key='synthetic-not-a-key')
assert callable(c.aio.live.connect)
assert types.Blob(data=b'\0\0',mime_type='audio/pcm;rate=16000').data==b'\0\0'
c.close()
for name in ['workspace_state.py','langgraph_review.py','mcp_client.py']:
 run=subprocess.run([sys.executable,str(EXAMPLES/name)],capture_output=True,text=True,timeout=30)
 assert run.returncode==0,(name,run.stderr)
 print(name+': '+run.stdout.strip())
print('PASS: real FastAPI WebSocket auth/PCM/stop/error, PDF/DOCX/VTT parsing, signature/review boundaries, installed Live SDK contract, LangGraph interrupt/resume and real MCP stdio round trip. No live provider quality claim.')
# Exercise the embedding HTTP contract with controlled transport responses.
import httpx
from embedding_adapter import embed
async def embedding_checks():
 def handler(request):
  payload=json.loads(request.content);assert payload['truncate'] is False
  assert payload['input']==['Synthetic policy']
  return httpx.Response(200,json={'model':'embedding-fixture','embeddings':[[1.,0.,0.]],'prompt_eval_count':2})
 transport=httpx.MockTransport(handler)
 result=await embed(['Synthetic policy'],'embedding-fixture',3,transport=transport)
 assert result['input_tokens']==2 and len(result['vectors'])==1
 try:await embed(['Synthetic policy'],'embedding-fixture',2,transport=transport)
 except ValueError:pass
 else:raise AssertionError('dimension mismatch accepted')
asyncio.run(embedding_checks())
from workspace_connectors import Connectors
from unittest.mock import patch
async def connector_checks():
 requests=[];real_client=httpx.AsyncClient
 def handler(request):
  requests.append(request)
  if request.method=='GET':return httpx.Response(200,json={'id':'message-1','payload':{'headers':[{'name':'Subject','value':'Training'}]}})
  assert request.url.path.endswith('/drafts') and '/send' not in request.url.path
  import base64
  message=json.loads(request.content)['message']['raw']
  decoded=base64.urlsafe_b64decode(message).decode();assert 'Synthetic draft' in decoded
  return httpx.Response(200,json={'id':'draft-1'})
 client=Connectors('synthetic-slack','synthetic-google',['test-channel'],'test-calendar')
 with patch('workspace_connectors.httpx.AsyncClient',side_effect=lambda **kw:real_client(transport=httpx.MockTransport(handler),**kw)):
  assert (await client.gmail_metadata('message-1'))['id']=='message-1'
  payload={'operation':'op-1','to':'learner@example.invalid','subject':'Training','body':'Synthetic draft'}
  approval={'owner':'training','hash':payload_hash(payload),'expires':time.time()+60}
  result=await client.gmail_reviewed_draft('training','op-1',payload['to'],payload['subject'],payload['body'],approval)
  assert result['state']=='draft_created' and len(requests)==2
  try:await client.gmail_reviewed_draft('training','op-1','changed@example.invalid',payload['subject'],payload['body'],approval)
  except PermissionError:pass
  else:raise AssertionError('changed recipient accepted')
  assert len(requests)==2
import time
asyncio.run(connector_checks())
print('PASS: mocked embedding HTTP batches/dimension rejection and reviewed Gmail draft-only request; changed-recipient effect denied before HTTP.')
